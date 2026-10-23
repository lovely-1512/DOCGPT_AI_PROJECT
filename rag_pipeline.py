"""DocGPT AI — Evidence-first backend: load → chunk → embed → RRF → rerank → Groq LLM."""
import os, re, hashlib, time
import numpy as np, pandas as pd, faiss
from groq import Groq
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer, CrossEncoder
from pypdf import PdfReader
from pptx import Presentation
from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
RERANKER_MODEL  = "cross-encoder/ms-marco-MiniLM-L-6-v2"
GROQ_MODEL      = "openai/gpt-oss-120b"
GROQ_API_KEY    = os.environ.get("GROQ_API_KEY")
if not GROQ_API_KEY: raise RuntimeError("GROQ_API_KEY missing")
groq_client = Groq(api_key=GROQ_API_KEY)
RRF_K, DEFAULT_FINAL_K, DEFAULT_CANDIDATE_K, MAX_CONTEXT_CHARS = 60, 5, 12, 12000

print("Loading models...")
embedding_model = SentenceTransformer(EMBEDDING_MODEL)
reranker = CrossEncoder(RERANKER_MODEL, max_length=512)
print("Models loaded with Groq Cloud.")

# ---------- utils ----------
def create_document_id(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""): h.update(b)
    return h.hexdigest()[:16]

def clean_text(t):
    t = re.sub(r"\s+", " ", t)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", t).strip()

def clean_llm_output(t):
    if not t: return ""
    t = str(t).strip()
    if "<think>" in t:
        parts = t.split("</think>", 1)
        t = parts[1].strip() if len(parts) == 2 else t.split("<think>", 1)[0].strip()
    t = re.sub(r"<think>.*?</think>", "", t, flags=re.DOTALL | re.IGNORECASE)
    for p in [r"^\s*Thinking Process.*$", r"^\s*Constraint Checklist.*$",
              r"^\s*Analyze the Request.*$", r"^\s*Analyze the Evidence.*$",
              r"^\s*Confidence Score.*$", r"^\s*Return ONLY.*$"]:
        t = re.sub(p, "", t, flags=re.IGNORECASE | re.MULTILINE)
    return t.strip()

def tokenize_bm25(t): return re.findall(r"\b\w+\b", t.lower())

# ---------- loaders ----------
def read_pdf(p):
    did, recs = create_document_id(p), []
    for i, pg in enumerate(PdfReader(p).pages, 1):
        t = clean_text(pg.extract_text() or "")
        if t: recs.append(dict(document_id=did, source=os.path.basename(p), file_type="PDF",
            location_type="page", location=f"Page {i}", page=i, slide=None, text=t))
    return recs

def read_pptx(p):
    did, recs = create_document_id(p), []
    for i, s in enumerate(Presentation(p).slides, 1):
        t = clean_text("\n".join(sh.text.strip() for sh in s.shapes if hasattr(sh,"text") and sh.text.strip()))
        if t: recs.append(dict(document_id=did, source=os.path.basename(p), file_type="PPTX",
            location_type="slide", location=f"Slide {i}", page=None, slide=i, text=t))
    return recs

def read_docx(p):
    did, recs, n = create_document_id(p), [], 0
    for para in Document(p).paragraphs:
        t = clean_text(para.text)
        if t:
            n += 1
            recs.append(dict(document_id=did, source=os.path.basename(p), file_type="DOCX",
                location_type="paragraph", location=f"Paragraph {n}", page=None, slide=None, text=t))
    return recs

def read_txt(p):
    did = create_document_id(p); t = ""
    for enc in ["utf-8","utf-8-sig","cp1252","latin-1"]:
        try: t = open(p, encoding=enc).read(); break
        except UnicodeDecodeError: continue
    t = clean_text(t)
    return [dict(document_id=did, source=os.path.basename(p), file_type="TXT",
        location_type="document", location="Full document", page=None, slide=None, text=t)] if t else []

def read_tabular(p):
    did = create_document_id(p); ext = p.lower().rsplit(".",1)[-1]
    try: df = pd.read_csv(p) if ext=="csv" else pd.read_excel(p)
    except UnicodeDecodeError: df = pd.read_csv(p, encoding="latin-1")
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    out = []
    for rn, row in df.fillna("").iterrows():
        t = clean_text(" | ".join(f"{c}: {row[c]}" for c in df.columns if str(row[c]).strip()))
        if t: out.append(dict(document_id=did, source=os.path.basename(p), file_type=ext.upper(),
            location_type="row", location=f"Row {rn+2}", page=None, slide=None, text=t))
    return out

def read_xlsx(p):
    did, out = create_document_id(p), []
    for sh in pd.ExcelFile(p).sheet_names:
        df = pd.read_excel(p, sheet_name=sh).dropna(axis=0, how="all").dropna(axis=1, how="all")
        for rn, row in df.fillna("").iterrows():
            t = clean_text(" | ".join(f"{c}: {row[c]}" for c in df.columns if str(row[c]).strip()))
            if t: out.append(dict(document_id=did, source=os.path.basename(p), file_type="XLSX",
                location_type="sheet_row", location=f"Sheet: {sh} • Row: {rn+2}",
                page=None, slide=None, text=t))
    return out

def load_document(p):
    ext = p.lower().rsplit(".", 1)[-1]
    return {"pdf":read_pdf,"pptx":read_pptx,"docx":read_docx,"txt":read_txt,
            "csv":read_tabular,"xlsx":read_xlsx}.get(ext, lambda _: [])(p)

# ---------- engine ----------
class DocGPTEngine:
    def __init__(self):
        self.chunk_df = pd.DataFrame(); self.index = None; self.bm25 = None
        self.document_units = []; self.embeddings = None
        self.ready = False; self.documents = []; self.total_chunks = 0

    def build_knowledge_base(self, file_paths):
        t0 = time.time(); units = []
        for p in file_paths: units.extend(load_document(p))
        if not units:
            self.ready = False
            return {"success":False, "message":"No readable text.", "documents":[], "chunks":0}
        self.document_units = units
        splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=120,
            separators=["\n\n","\n",". ","? ","! ","; "," ",""])
        recs, cid = [], 0
        for u in units:
            for n, chunk in enumerate(splitter.split_text(u["text"]), 1):
                recs.append({**{k:u[k] for k in ["document_id","source","file_type",
                    "location_type","location","page","slide"]},
                    "chunk_id":cid, "chunk_number":n, "text":chunk,
                    "indexed_at":time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())})
                cid += 1
        self.chunk_df = pd.DataFrame(recs)
        if self.chunk_df.empty:
            self.ready = False
            return {"success":False, "message":"No usable chunks.", "documents":[], "chunks":0}
        texts = self.chunk_df["text"].tolist()
        self.embeddings = embedding_model.encode(texts, batch_size=32,
            show_progress_bar=False, normalize_embeddings=True).astype("float32")
        self.index = faiss.IndexFlatIP(self.embeddings.shape[1]); self.index.add(self.embeddings)
        self.bm25 = BM25Okapi([tokenize_bm25(t) for t in texts])
        self.documents = sorted(self.chunk_df["source"].unique().tolist())
        self.total_chunks = len(self.chunk_df); self.ready = True
        return {"success":True, "message":"Knowledge base created.",
                "documents":self.documents, "chunks":self.total_chunks,
                "document_units":len(self.document_units),
                "processing_time":round(time.time()-t0, 2)}

    def hybrid_retrieve(self, query, candidate_k=DEFAULT_CANDIDATE_K):
        if not self.ready: return []
        qe = embedding_model.encode([query], normalize_embeddings=True).astype("float32")
        _, sem = self.index.search(qe, min(candidate_k, len(self.chunk_df)))
        sem = sem[0].tolist()
        kw = np.argsort(self.bm25.get_scores(tokenize_bm25(query)))[::-1][:candidate_k].tolist()
        rrf = {}
        for rank, i in enumerate(sem, 1): rrf[i] = rrf.get(i,0) + 1/(RRF_K+rank)
        for rank, i in enumerate(kw, 1):  rrf[i] = rrf.get(i,0) + 1/(RRF_K+rank)
        out = []
        for i, s in sorted(rrf.items(), key=lambda x: x[1], reverse=True):
            r = self.chunk_df.iloc[i]
            out.append({"chunk_id":int(r["chunk_id"]),"document_id":r["document_id"],
                "source":r["source"],"file_type":r["file_type"],"location":r["location"],
                "page":r["page"],"slide":r["slide"],"chunk_number":int(r["chunk_number"]),
                "text":r["text"],"rrf_score":float(s)})
        return out

    def rerank(self, query, candidates, final_k=DEFAULT_FINAL_K):
        if not candidates: return []
        scores = reranker.predict([(query, c["text"]) for c in candidates])
        for c, s in zip(candidates, scores): c["reranker_score"] = float(s)
        return sorted(candidates, key=lambda x: x["reranker_score"], reverse=True)[:final_k]

    def calculate_retrieval_confidence(self, evidence):
        if not evidence:
            return {"score":0.0, "label":"Low", "type":"heuristic"}
        scores = np.array([e.get("reranker_score", e.get("score",0.0)) for e in evidence])
        top_sig = 1/(1+np.exp(-np.clip(scores[0], -10, 10)))
        ev_sig = min(len(evidence)/5.0, 1.0)
        mgn = scores[0]-scores[1] if len(scores) >= 2 else 0.0
        mgn_sig = 1/(1+np.exp(-np.clip(mgn, -10, 10)))
        c = float(np.clip(0.5*top_sig + 0.3*ev_sig + 0.2*mgn_sig, 0, 1))
        return {"score":c, "label":"High" if c >= 0.75 else "Medium" if c >= 0.5 else "Low",
                "type":"heuristic"}

    def build_evidence_context(self, evidence):
        blocks, cur = [], 0
        for i, e in enumerate(evidence, 1):
            b = f"[EVIDENCE {i}]\nSource: {e['source']}\nLocation: {e['location']}\nContent:\n{e['text']}\n"
            if cur + len(b) > MAX_CONTEXT_CHARS: break
            blocks.append(b); cur += len(b)
        return "\n".join(blocks)

    def ask_rag(self, question, conversation_history=None, candidate_k=DEFAULT_CANDIDATE_K,
                final_k=DEFAULT_FINAL_K, max_new_tokens=300):
        t0 = time.time()
        if not self.ready:
            return {"answer":"Upload files first.","citations":[],"evidence":[],
                    "confidence":{"score":0.0,"label":"Low","type":"heuristic"},"latency":0.0}
        evidence = self.rerank(question, self.hybrid_retrieve(question, candidate_k), final_k)
        if not evidence:
            return {"answer":"No evidence found.","citations":[],"evidence":[],
                    "confidence":{"score":0.0,"label":"Low","type":"heuristic"},
                    "latency":round(time.time()-t0, 3)}
        ctx = self.build_evidence_context(evidence)
        hist = ""
        if conversation_history:
            hist = "\n".join(f"{m.get('role','')}: {m.get('content','')}"
                             for m in conversation_history[-6:] if m.get("content"))
        system = ("You are DocGPT AI. Answer using ONLY the supplied evidence. "
                  "Do not use outside knowledge. If unsupported, say "
                  "'I could not find sufficient information in the uploaded documents.' "
                  "Answer naturally. Use bullets for multiple facts, **bold** for key terms. "
                  "Never show reasoning, thinking, or <think> tags. Do not create a sources section.")
        user = (f"DOCUMENT CONTEXT:\n\n{ctx}\n\nUSER QUESTION:\n{question}\n\n"
                f"CONVERSATION:\n{hist}\n\nAnswer using only the document context.")
        try:
            resp = groq_client.chat.completions.create(
                model=GROQ_MODEL,
                messages=[{"role":"system","content":system},{"role":"user","content":user}],
                temperature=0.7, max_tokens=max_new_tokens,
                reasoning_effort="low", reasoning_format="hidden")
            answer = clean_llm_output(resp.choices[0].message.content or "")
        except Exception as e:
            answer = f"Groq API Error: {e}"
        citations = [{"evidence_id":i,"source":r["source"],"location":r["location"],
                      "page":r["page"],"slide":r["slide"],
                      "score":r["reranker_score"]} for i, r in enumerate(evidence, 1)]
        return {"answer":answer, "citations":citations, "evidence":evidence,
                "confidence":self.calculate_retrieval_confidence(evidence),
                "latency":round(time.time()-t0, 3)}

    def generate_study_material(self, topic, study_type, max_new_tokens=700):
        if not self.ready:
            return {"success":False,"content":"Please upload and process documents first.",
                    "evidence":[],"confidence":{"score":0.0,"label":"Low","type":"heuristic"}}
        evidence = self.rerank(topic, self.hybrid_retrieve(topic, 15), 6)
        if not evidence:
            return {"success":False,"content":"Insufficient evidence for this topic.",
                    "evidence":[],"confidence":{"score":0.0,"label":"Low","type":"heuristic"}}
        ctx = self.build_evidence_context(evidence)
        if study_type == "Viva Questions":
            inst = ("Create exactly 10 technical viva questions with short answers.\n"
                    "Return ONLY final Q&A. Format:\n### Question N\n**Question:** ...\n"
                    "**Answer:** ...\n**Source:** [Evidence N]\n"
                    "Use ONLY the evidence. No outside facts. Start with ### Question 1.")
        elif study_type.startswith("Multiple Choice Questions"):
            m = re.search(r"\((\d+)\s*questions?\)", study_type, re.IGNORECASE)
            n = int(m.group(1)) if m else 5
            inst = (f"Create exactly {n} MCQs. Return ONLY quiz. Format:\n"
                    f"1. Question\nA) ...\nB) ...\nC) ...\nD) ...\nCorrect Answer: A\n"
                    f"Exactly {n} questions. Four options each. One correct. No extra sections.")
        elif study_type == "Flashcards":
            inst = ("Create exactly 5 flashcards. Format:\n### Flashcard N\n**Question:** ...\n"
                    "**Answer:** ...\n**Key Term:** **term**\n**Source:** [Evidence N]\n"
                    "Use ONLY the evidence.")
        else:
            inst = ("Create structured revision notes:\n### Topic Overview\n### Key Concepts\n"
                    "### Important Definitions\n### Important Facts\n### Quick Revision\n"
                    "Use ONLY the evidence.")
        system = ("You are DocGPT AI Study Assistant. Use ONLY supplied evidence. "
                  "Never invent facts/citations. Return ONLY final content. No reasoning, "
                  "no analysis, no chain-of-thought. Use ### headings and **bold** key terms.")
        user = (f"Topic: {topic}\nStudy Mode: {study_type}\nTask:\n{inst}\n\n"
                f"EVIDENCE:\n{ctx}\n\nGenerate final study material now.")
        try:
            resp = groq_client.chat.completions.create(
                model=GROQ_MODEL,
                messages=[{"role":"system","content":system},{"role":"user","content":user}],
                temperature=0.0, max_tokens=max_new_tokens,
                reasoning_effort="low", reasoning_format="hidden")
            content = clean_llm_output(resp.choices[0].message.content or "")
        except Exception as e:
            content = f"Groq API Error: {e}"
        return {"success":True, "content":content, "evidence":evidence,
                "confidence":self.calculate_retrieval_confidence(evidence)}

        def document_health(self):
         if self.chunk_df.empty: return pd.DataFrame()
         rows = []
         for src, g in self.chunk_df.groupby("source"):
            rows.append({"Document":src, "File Type":g["file_type"].iloc[0],
                "Chunks":len(g), "Locations":g["location"].nunique(),
                "Characters":int(g["text"].str.len().sum()), "Status":"Indexed"})
         return pd.DataFrame(rows)

    def compare_documents(self, document_a, document_b, question,
                          top_k=8, max_new_tokens=500):
        t0 = time.time()
        if not self.ready:
            return {"success":False, "answer":"Upload and process documents first.",
                    "evidence":[], "confidence":{"score":0.0,"label":"Low","type":"heuristic"},
                    "documents":{"A":document_a,"B":document_b}}
        a = self.chunk_df[self.chunk_df["source"] == document_a].copy()
        b = self.chunk_df[self.chunk_df["source"] == document_b].copy()
        if a.empty or b.empty:
            return {"success":False, "answer":"Insufficient content in one or both documents.",
                    "evidence":[], "confidence":{"score":0.0,"label":"Low","type":"heuristic"},
                    "documents":{"A":document_a,"B":document_b}}
        a["reranker_score"] = reranker.predict([(question, t) for t in a["text"]])
        b["reranker_score"] = reranker.predict([(question, t) for t in b["text"]])
        a = a.sort_values("reranker_score", ascending=False)
        b = b.sort_values("reranker_score", ascending=False)
        k = max(2, top_k // 2)
        combined = pd.concat([a.head(k), b.head(k)], ignore_index=True) \
                     .sort_values("reranker_score", ascending=False)
        results = [{"source":r["source"], "location":r["location"], "page":r["page"],
                    "slide":r["slide"], "text":r["text"],
                    "reranker_score":float(r["reranker_score"]),
                    "score":float(r["reranker_score"])}
                   for _, r in combined.iterrows()]
        blocks, cur, n = [], 0, 1
        for item in results:
            label = "DOCUMENT A" if item["source"] == document_a else "DOCUMENT B"
            blk = (f"[EVIDENCE {n}]\nDocument: {label}\nSource: {item['source']}\n"
                   f"Location: {item['location']}\nContent:\n{item['text']}\n")
            if cur + len(blk) > MAX_CONTEXT_CHARS: break
            blocks.append(blk); cur += len(blk); n += 1
        ctx = "\n".join(blocks)
        system = ("You are DocGPT AI Comparison Assistant. Compare two documents using "
                  "ONLY supplied evidence. No invented differences/similarities. "
                  "Cite [Evidence N]. Use ### headings, bullets, and a Markdown table when useful. "
                  "Structure: ### Comparison Summary / ### Key Differences / ### Similarities / "
                  "### What Changed? / ### Conclusion. Bold important terms.")
        user = (f"DOCUMENT A: {document_a}\nDOCUMENT B: {document_b}\n"
                f"QUESTION: {question}\n\nEVIDENCE:\n{ctx}\n\nCompare A and B using only evidence.")
        try:
            resp = groq_client.chat.completions.create(
                model=GROQ_MODEL,
                messages=[{"role":"system","content":system},{"role":"user","content":user}],
                temperature=0.7, max_tokens=max_new_tokens,
                reasoning_effort="low", reasoning_format="hidden")
            answer = clean_llm_output(resp.choices[0].message.content or "")
        except Exception as e:
            answer = f"Groq API Comparison Error: {e}"
        citations = [{"evidence_id":i, "source":r["source"], "location":r["location"],
                      "score":r["reranker_score"]} for i, r in enumerate(results, 1)]
        return {"success":True, "answer":answer, "evidence":results, "citations":citations,
                "confidence":self.calculate_retrieval_confidence(results),
                "documents":{"A":document_a,"B":document_b},
                "latency":round(time.time()-t0, 3)}


def create_engine():
    return DocGPTEngine()
